"""
Script to generate the sample FlowForge Excel factory template.
Saved to examples/factory_data_template.xlsx
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_sample_template():
    wb = openpyxl.Workbook()

    # --- Sheet 1: Machines ---
    ws_machines = wb.active
    ws_machines.title = "Machines"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="F8FAFC")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    machine_headers = ["machine_id", "machine_name", "status", "energy_kwh_per_hour", "available_from", "available_until"]
    ws_machines.append(machine_headers)

    machines_data = [
        ["M1", "CNC Milling Center 01", "available", 8.0,  "08:00", "18:00"],
        ["M2", "CNC Lathe Heavy Duty 02", "available", 10.5, "08:00", "18:00"],
        ["M3", "High Precision CNC 03", "available", 12.0, "08:00", "18:00"],
        ["M4", "Laser Cutting Cell 04", "available", 7.5,  "08:00", "18:00"],
        ["M5", "Robotic Assembly Cell 05", "available", 7.0,  "08:00", "18:00"],
        ["M6", "Automated Inspection 06", "available", 9.0,  "08:00", "18:00"],
    ]

    for row in machines_data:
        ws_machines.append(row)

    # --- Sheet 2: Jobs ---
    ws_jobs = wb.create_sheet(title="Jobs")
    job_headers = ["job_id", "job_name", "processing_time", "eligible_machines", "priority", "deadline", "quantity", "status"]
    ws_jobs.append(job_headers)

    jobs_data = [
        ["J1",  "Engine Housing Block",    45, "M1,M2,M3,M4,M5,M6", "High",    200, 1, "pending"],
        ["J2",  "Drive Shaft Precision",   30, "M1,M2,M3,M4,M5,M6", "High",    120, 1, "pending"],
        ["J3",  "Turbine Mounting Plate",  60, "M1,M2,M3,M4,M5,M6", "Minimal", 250, 1, "pending"],
        ["J4",  "Piston Rod Assembly",     35, "M1,M2,M3,M4,M5,M6", "Medium",  150, 1, "pending"],
        ["J5",  "Gearbox Main Casing",     50, "M1,M2,M3,M4,M5,M6", "High",    180, 1, "pending"],
        ["J6",  "Sensor Bracket Light",    25, "M1,M2,M3,M4,M5,M6", "Critical",100, 1, "pending"],
        ["J7",  "Exhaust Manifold",        55, "M1,M2,M3,M4,M5,M6", "Medium",  220, 1, "pending"],
        ["J8",  "Hydraulic Valve Body",    40, "M1,M2,M3,M4,M5,M6", "High",    160, 1, "pending"],
        ["J9",  "Fastener Pin Batch",      20, "M1,M2,M3,M4,M5,M6", "Critical", 90, 1, "pending"],
        ["J10", "Main Chassis Frame",      65, "M1,M2,M3,M4,M5,M6", "Minimal", 300, 1, "pending"],
        ["J11", "Fuel Injection Nozzle",   30, "M1,M2,M3,M4,M5,M6", "Medium",  140, 1, "pending"],
        ["J12", "Brake Rotor Disc",        45, "M1,M2,M3,M4,M5,M6", "Medium",  260, 1, "pending"],
        ["J13", "Bearing Retainer Ring",   35, "M1,M2,M4,M5",       "High",    175, 1, "pending"],
        ["J14", "Cooling Fan Shroud",      40, "M2,M3,M5,M6",       "Medium",  210, 1, "pending"],
        ["J15", "Steering Column Shaft",   50, "M1,M3,M4,M6",       "High",    230, 1, "pending"],
        ["J16", "Control Panel Cover",     25, "M1,M2,M4,M5",       "Critical",110, 1, "pending"],
    ]

    for row in jobs_data:
        ws_jobs.append(row)

    # --- Sheet 3: Disruptions (Optional) ---
    ws_disruptions = wb.create_sheet(title="Disruptions")
    disruption_headers = ["disruption_id", "type", "target_id", "timestamp", "value"]
    ws_disruptions.append(disruption_headers)

    disruptions_data = [
        ["D1", "machine_failure", "M3", "10:30", None],
        ["D2", "urgent_job", "J17", "11:00", 30],
        ["D3", "deadline_change", "J7", "11:15", 130],
    ]

    for row in disruptions_data:
        ws_disruptions.append(row)

    # Styling sheets
    for ws in [ws_machines, ws_jobs, ws_disruptions]:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    os.makedirs("examples", exist_ok=True)
    out_path = os.path.join("examples", "factory_data_template.xlsx")
    wb.save(out_path)
    print(f"Generated sample template at {out_path}")


if __name__ == "__main__":
    generate_sample_template()
